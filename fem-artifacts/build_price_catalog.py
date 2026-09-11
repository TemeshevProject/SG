#!/usr/bin/env python3
"""Build consolidated CAPEX/OPEX artifact catalog from all historical FEM workbooks."""

from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

BASE = Path(__file__).resolve().parent
HIST_DIR = BASE / "historical"
OUTPUT = BASE / "fem_all_artifacts_prices.xlsx"
META = BASE / "price_catalog_meta.json"

SKIP_SHEETS = re.compile(
    r"(титул|titul|wacc|enpv|ддс|опиу|opiu|структур|index|пок\s*пр|пок\s*чп|"
    r"оценка\s+выгод|оценка\s+эфф|баланс\s+риск|чувствитель|график\s+выплат|"
    r"^займ$|макр|бюдж|база|^итог$|summ|inputs|composition|maintenance|"
    r"depreciation|financing|сэф|сээ|бип|сравнен|расчет\s+тариф|расчет\s+займа|"
    r"расчет\s+поступлен|площ$|ст\s+об|экон\s+пок|оц\s+выг|л ist4|list4|"
    r"действующ|расчет\s+ву|capEx\s+модерн|займ\s+модерн)",
    re.I,
)

CAPEX_SHEETS = re.compile(
    r"(инвест|capex|3\.capex|4\.capex|для\s*capex|2\.1\.1\.|rcou|"
    r"3\.\s*инвест|4\.\s*capex|capEx|с\.m\b|сm\b|расшифров)",
    re.I,
)

OPEX_SHEETS = re.compile(
    r"(экспл|opex|обсл|фот|для\s*opex|для\s*ил|зп\b|4\.\s*opex|7\.\s*opex|"
    r"8\.\s*opex|5\.\s*экспл|6\.\s*обсл|расчет\s+фот|расчет\s+экспл|"
    r"стоимости\s+услуг|4\.1\.фот|cost\b|2\.2\.maintenance|штраф)",
    re.I,
)

NAME_HEADERS = (
    "наименование",
    "должность",
    "наименование ",
    "наименование работ",
    "наименование раздела",
)

PRICE_HEADERS = (
    ("итоговая стоимость", 110),
    ("себестоимость за ед.с ндс", 100),
    ("себестоимость за ед", 100),
    ("цена за ед", 100),
    ('зп "к начислению"', 90),
    ("зп к начислению", 90),
    ("оклад с учетом налогов", 90),
    ("зп на руки", 85),
    ("цена в мес", 80),
    ("цена в мес, тыс", 75),
    ("затраты в год", 70),
    ("стоимость работ", 68),
    ("стоимость оборудования", 67),
    ("итого себестоимость", 65),
)


def is_qty_header(h: str) -> bool:
    return ("кол" in h or "кол-во" in h) and "себест" not in h and "затрат" not in h


def pick_price_columns(headers: dict[int, str]) -> list[tuple[int, int]]:
    scored = []
    for c, h in headers.items():
        if is_qty_header(h):
            continue
        for needle, score in PRICE_HEADERS:
            if needle in h:
                scored.append((score, c))
                break
    scored.sort(reverse=True)
    return scored

EXCLUDE_NAME = re.compile(
    r"(^итого|^всего|^амортизац|^остаточн|^№\s*п/п$|"
    r"^наименование\s+показател|^наименование\s+работ$|^фот,?\s*тенге$|"
    r"инвестиционные\s+затрат|операционные\s+затрат|"
    r"^\d+[\.\)]?\s*$|^none$|^\d{4}[\.\-]\d+$)",
    re.I,
)


def norm(text) -> str:
    if text is None:
        return ""
    s = str(text).strip().lower().replace("ё", "е")
    s = re.sub(r"\s+", " ", s)
    return s


def to_float(val):
    if val is None:
        return None
    if isinstance(val, (int, float)):
        if isinstance(val, float) and (val != val):  # NaN
            return None
        return float(val)
    s = str(val).strip().replace(" ", "").replace(",", ".")
    if not s or s in {"-", "#REF!", "#N/A"}:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def classify_from_headers(headers: dict[int, str]) -> str | None:
    joined = " ".join(headers.values())
    if "себестоимость за ед" in joined or "итого себестоимость" in joined:
        return "CAPEX"
    if any(x in joined for x in ("цена в мес", "затраты в год", "зп на руки", "зп к начислению", "оклад с учетом")):
        return "OPEX"
    return None


def classify_sheet(sheet_name: str, ws, headers: dict[int, str] | None = None) -> str | None:
    sn = norm(sheet_name)
    if SKIP_SHEETS.search(sheet_name):
        return None
    if CAPEX_SHEETS.search(sheet_name):
        return "CAPEX"
    if OPEX_SHEETS.search(sheet_name):
        return "OPEX"
    if headers:
        h_kind = classify_from_headers(headers)
        if h_kind:
            return h_kind
    blob = ""
    for r in range(1, 8):
        for c in range(1, 6):
            v = ws.cell(r, c).value
            if v:
                blob += " " + str(v).lower()
    if "инвестицион" in blob and "затрат" in blob:
        return "CAPEX"
    if "операцион" in blob or "эксплуатац" in blob:
        return "OPEX"
    if "наименование" in blob and ("себестоимость" in blob or "инвест" in blob):
        return "CAPEX"
    if "наименование" in blob and ("цена в мес" in blob or "затраты в год" in blob):
        return "OPEX"
    return None


def find_table_header(ws, max_row=40):
    best = None
    best_headers = None
    for r in range(1, max_row + 1):
        headers = {}
        for c in range(1, min(ws.max_column, 25) + 1):
            h = norm(ws.cell(r, c).value)
            if not h:
                continue
            headers[c] = h
        if not headers:
            continue
        name_col = None
        for c, h in headers.items():
            if any(h.startswith(n.strip()) or n.strip() in h for n in NAME_HEADERS):
                name_col = c
                break
        if not name_col:
            continue
        price_col = None
        price_score = -1
        for c, h in headers.items():
            if is_qty_header(h):
                continue
            for needle, score in PRICE_HEADERS:
                if needle in h and score > price_score:
                    price_col = c
                    price_score = score
        if price_col:
            best = (r, name_col, price_col)
            best_headers = headers
    if best:
        return (*best, best_headers)
    return None


def is_valid_name(name: str) -> bool:
    if not name or len(name.strip()) < 3:
        return False
    n = norm(name)
    if EXCLUDE_NAME.search(n):
        return False
    if re.fullmatch(r"[\d\.,\s]+", n):
        return False
    return True


def row_price(
    ws,
    r: int,
    name_col: int,
    headers: dict[int, str],
    header_row: int,
    capex_opex: str,
    primary_col: int,
):
    candidates: list[float] = []
    for _, c in pick_price_columns(headers):
        v = to_float(ws.cell(r, c).value)
        if v is not None and v > 0:
            candidates.append(v)
    if not candidates:
        for c in range(name_col + 1, min(name_col + 12, ws.max_column + 1)):
            h = headers.get(c)
            if h is None:
                h = norm(ws.cell(header_row, c).value)
            if is_qty_header(h):
                continue
            v = to_float(ws.cell(r, c).value)
            if v is not None and v > 0:
                candidates.append(v)
    if not candidates:
        return None

    for c, h in headers.items():
        if "итогов" in h and "стоим" in h:
            total = to_float(ws.cell(r, c).value)
            if total is not None and total > 0:
                return total

    primary = to_float(ws.cell(r, primary_col).value)
    if primary is not None and primary > 0 and not is_qty_header(headers.get(primary_col, "")):
        if capex_opex == "OPEX" or primary >= 1000:
            return primary

    big = [v for v in candidates if v >= 1000]
    if big:
        return max(big)
    return max(candidates)


def extract_sheet(ws, sheet_name: str, fem_name: str, capex_opex: str | None) -> list[dict]:
    header = find_table_header(ws)
    if not header:
        return []
    header_row, name_col, price_col, headers = header
    if not capex_opex:
        capex_opex = classify_from_headers(headers) or classify_sheet(sheet_name, ws, headers)
    if not capex_opex:
        return []
    out = []
    for r in range(header_row + 1, ws.max_row + 1):
        name = ws.cell(r, name_col).value
        if not name:
            continue
        name_s = str(name).strip()
        if not is_valid_name(name_s):
            continue
        price = row_price(ws, r, name_col, headers, header_row, capex_opex, price_col)
        out.append(
            {
                "type": capex_opex,
                "name": name_s,
                "price": price,
                "source": f"{fem_name} :: {sheet_name}",
            }
        )
    return out


def extract_workbook(path: Path) -> list[dict]:
    items = []
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:
        print(f"Skip {path.name}: {exc}")
        return items

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if SKIP_SHEETS.search(sheet_name):
            continue
        kind = classify_sheet(sheet_name, ws)
        items.extend(extract_sheet(ws, sheet_name, path.name, kind))
    wb.close()
    return items


def dedupe_max_price(items: list[dict]) -> list[dict]:
    buckets: dict[str, list[dict]] = {}
    for it in items:
        key = (it["type"], norm(it["name"]))
        buckets.setdefault(key, []).append(it)

    result = []
    for key, group in buckets.items():
        with_price = [g for g in group if g["price"] is not None and g["price"] > 0]
        if with_price:
            best = max(with_price, key=lambda x: x["price"])
        else:
            best = group[0]
        result.append(best)
    result.sort(key=lambda x: (x["type"], norm(x["name"])))
    return result


def write_output(rows: list[dict]):
    wb = openpyxl.Workbook()
    headers = ["CAPEX/OPEX", "Наименование артефакта", "Цена", "Источник"]

    def fill_sheet(ws, data):
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9E1F2")
        for row in data:
            ws.append([row["type"], row["name"], row["price"], row["source"]])
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 70
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 55

    ws_all = wb.active
    ws_all.title = "Все артефакты"
    fill_sheet(ws_all, rows)

    for title, flt in [("CAPEX", "CAPEX"), ("OPEX", "OPEX")]:
        ws = wb.create_sheet(title)
        fill_sheet(ws, [r for r in rows if r["type"] == flt])

    wb.save(OUTPUT)


def main():
    if not HIST_DIR.exists():
        raise SystemExit(f"Missing {HIST_DIR}")

    raw = []
    for path in sorted(HIST_DIR.glob("*.xlsx")):
        extracted = extract_workbook(path)
        print(f"{path.name}: {len(extracted)} rows")
        raw.extend(extracted)

    deduped = dedupe_max_price(raw)
    write_output(deduped)

    meta = {
        "historical_files": len(list(HIST_DIR.glob("*.xlsx"))),
        "raw_rows": len(raw),
        "unique_artifacts": len(deduped),
        "capex": sum(1 for r in deduped if r["type"] == "CAPEX"),
        "opex": sum(1 for r in deduped if r["type"] == "OPEX"),
        "with_price": sum(1 for r in deduped if r["price"] is not None and r["price"] > 0),
        "output": str(OUTPUT.name),
    }
    META.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(meta, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
