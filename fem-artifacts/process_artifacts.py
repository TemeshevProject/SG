#!/usr/bin/env python3
"""Compare historical FEM artifacts with main registry and produce updated workbook."""

import json
import re
import shutil
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
MAIN = BASE / "main_artifacts.xlsx"
HIST_DIR = BASE / "historical"
OUTPUT = BASE / "main_artifacts_updated.xlsx"
REPORT = BASE / "comparison_report.json"

CAPEX_SHEETS = re.compile(
    r"(инвест|capex|3\.capex|4\.capex|3\. инвест|4\. capex)",
    re.I,
)
OPEX_SHEETS = re.compile(
    r"(экспл|обсл|opex|5\.\s*экспл|6\.\s*обсл|7\.\s*opex|8\.\s*opex|4\.\s*opex)",
    re.I,
)

CAPEX_CATEGORIES = {
    "аппаратно-программные комплексы",
    "камеры общественной безопасности",
    "центр обработки и анализа данных",
    "стоимость монтажных работ",
    "оснащение рабочих мест ситуационного центра",
    "сопутствующее оборудование",
    "расходы на приобретение оборудования",
    "расходы на монтажные работы",
    "операционные затраты на этапе строительства",
    "3. см",
    "4. см",
    "строительно-монтаж",
}

OPEX_CATEGORIES = {
    "аренда каналов связи",
    "обслуживание апк",
    "услуги сторонних организаций",
    "расчет затрат по организации канала",
    "расчет затрат, связанных с обслуживанием",
    "расчет затрат на аренду помещений",
    "экономические предположения",
    "операционные затраты",
    "командировочные",
    "заработная плата",
    "коммунальные",
}

EXCLUDE_NAMES = re.compile(
    r"(^итого|^всего|^амортизац|^остаточн|^наименование\s+показател|^наименование\s+работ$|"
    r"^монтажные\s+работы$|^№\s*п/п$|инвестиционные\s+затрат|операционные\s+затрат|"
    r"себестоимость|на\s+начало\s+года|начало\s+года|конец\s+года|"
    r"тыс\.\s*тенге$|^\d+[\.\)]?\s*$|^none$|^\d{4}[\.\-]\d+$|^\d+\.\d+$)",
    re.I,
)

HARDWARE_SKU = re.compile(
    r"(sfp\+?|dimm|ddr\d|fortigate|vmware|vsphere|esxi|"
    r"supermicro|dell\s+poweredge|hpe\s+proliant|"
    r"лицензи[яи]\s+(vmware|windows\s+server)|"
    r"трансивер|оптический\s+модуль|ssd\s+\d+gb|hdd\s+\d+tb)",
    re.I,
)

GENERIC_EXCLUDE = re.compile(
    r"^(стол|стул|тумба|товар|лицензии|компьютер|монитор|шкаф|диспенсер|"
    r"сопутствующее\s+оборудование|дорожная\s+безопасность|общественная\s+безопасность|"
    r"центр\s+обработки\s+данных|кол-во\s+узлов|наименование\s+работ|монтажные\s+работы|"
    r"шкаф\s+для|диспенсер\s+для)",
    re.I,
)

VENDOR_EXCLUDE = re.compile(r"\b(тоо|ао|ип\b|кар-тел|кселл|билайн|fortigite|fortigate)\b", re.I)

OPEX_WHITELIST = re.compile(
    r"(билеты|проживание|суточные|дт\s+для\s+дгу|охрана|канц\.?\s*товар|"
    r"хоз\.?\s*товар|клининг|гсм\s+для|затраты\s+на\s+электроэнергию|"
    r"непредвиденные\s+работы|интернет\s+в\s+цод|обслуживание\s+апк\s+патруль|"
    r"аренда\s+цод|коммунальные\s+расходы)",
    re.I,
)

SECTION_ONLY = re.compile(
    r"^(аппаратно-программные комплексы|камеры общественной безопасности|"
    r"центр обработки и анализа данных|аренда каналов связи|"
    r"услуги сторонних организаций|услуги по обслуживанию|"
    r"операционные затраты|инвестиционные затраты|"
    r"оснащение рабочих мест ситуационного центра)",
    re.I,
)


def should_add_missing(data):
    name = data["name"]
    n = norm(name)
    count = data["count"]
    typ = data["type"]

    if GENERIC_EXCLUDE.match(n):
        return False
    if VENDOR_EXCLUDE.search(name):
        return False
    if SECTION_ONLY.match(n) and len(n) < 80:
        return False
    if re.search(r"^настройка,?\s+пуск", n):
        return False
    if re.search(r"^(стоимость|затраты)\s+.*(начало|конец)\s+года", n):
        return False
    if re.fullmatch(r"[\d\.,\s]+", n):
        return False

    if typ == "OPEX":
        if OPEX_WHITELIST.search(name):
            return True
        if count >= 4 and len(n) >= 20:
            return True
        if is_important_product(name):
            return True
        return False

    if is_important_product(name):
        return True
    if count >= 3 and len(n) >= 18:
        return True
    if count >= 5:
        return True
    return False


def norm(text):
    if text is None:
        return ""
    s = str(text).strip().lower().replace("ё", "е")
    s = s.replace("«", '"').replace("»", '"').replace("'", '"')
    s = re.sub(r"\s+", " ", s)
    s = re.sub(r"[\"']", "", s)
    return s


def classify(category, sheet_type=None):
    c = norm(category)
    if sheet_type == "CAPEX":
        return "CAPEX"
    if sheet_type == "OPEX":
        return "OPEX"
    for key in CAPEX_CATEGORIES:
        if key in c:
            return "CAPEX"
    for key in OPEX_CATEGORIES:
        if key in c:
            return "OPEX"
    return "OPEX" if sheet_type == "OPEX" else "CAPEX"


def is_valid_artifact(name, category=None):
    if not name or not str(name).strip():
        return False
    n = norm(name)
    if len(n) < 4:
        return False
    if EXCLUDE_NAMES.search(n):
        return False
    if HARDWARE_SKU.search(n):
        return False
    if SECTION_ONLY.match(n) and len(n) < 60:
        return False
  # section headers often lack detail
    if n in {"наименование", "наименование показателя"}:
        return False
    return True


def find_header_row(ws, max_scan=15):
    for r in range(1, max_scan + 1):
        vals = [norm(ws.cell(r, c).value) for c in range(1, 8)]
        joined = " ".join(vals)
        if "наименование" in joined and ("п/п" in joined or "кол" in joined or "итого" in joined):
            return r
    return 3


def extract_from_sheet(ws, sheet_name, fem_name, sheet_type):
    header_row = find_header_row(ws)
    name_col = 2
    for c in range(1, 8):
        h = norm(ws.cell(header_row, c).value)
        if h.startswith("наименование"):
            name_col = c
            break

    current_category = ""
    results = []
    for r in range(header_row + 1, ws.max_row + 1):
        name = ws.cell(r, name_col).value
        if name is None:
            continue
        name_s = str(name).strip()
        n_norm = norm(name_s)

        # category header: numbered section without typical artifact detail
        first_cell = ws.cell(r, 1).value
        qty = ws.cell(r, 3).value if ws.max_column >= 3 else None
        is_section = False
        if isinstance(first_cell, str) and re.match(r"^\d+[\.\)]?\s*$", first_cell.strip()):
            is_section = False
        elif isinstance(first_cell, str) and re.match(r"^\d+[\.\)]", first_cell.strip()) and qty is None:
            is_section = True
        elif n_norm in SECTION_ONLY.pattern and qty is None:
            is_section = True

        if is_section or (SECTION_ONLY.match(n_norm) and not isinstance(first_cell, (int, float))):
            current_category = name_s
            continue

        if not is_valid_artifact(name_s, current_category):
            continue

        cat = current_category or sheet_name
        results.append(
            {
                "category": cat,
                "name": name_s,
                "type": sheet_type,
                "source": f"{fem_name}::{sheet_name}",
            }
        )
    return results


def extract_historical():
    all_items = []
    for path in sorted(HIST_DIR.glob("*.xlsx")):
        try:
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        except Exception as exc:
            print(f"Skip {path.name}: {exc}")
            continue
        for sheet_name in wb.sheetnames:
            st = None
            if CAPEX_SHEETS.search(sheet_name):
                st = "CAPEX"
            elif OPEX_SHEETS.search(sheet_name):
                st = "OPEX"
            else:
                continue
            ws = wb[sheet_name]
            all_items.extend(extract_from_sheet(ws, sheet_name, path.name, st))
        wb.close()
    return all_items


def load_main_artifacts():
    wb = openpyxl.load_workbook(MAIN, data_only=True)
    ws = wb["Артефакты"]
    rows = []
    for r in range(3, ws.max_row + 1):
        name = ws.cell(r, 3).value
        if not name:
            continue
        rows.append(
            {
                "row": r,
                "category": ws.cell(r, 2).value,
                "name": str(name).strip(),
                "unit": ws.cell(r, 4).value,
                "price": ws.cell(r, 5).value,
                "qty_supplier": ws.cell(r, 6).value,
                "stage": ws.cell(r, 7).value,
                "price_supplier": ws.cell(r, 8).value,
                "calc_holder": ws.cell(r, 9).value,
                "char_supplier": ws.cell(r, 10).value,
            }
        )
    return wb, rows


def aggregate_historical(items):
    agg = {}
    for it in items:
        key = norm(it["name"])
        if key not in agg:
            agg[key] = {
                "category": it["category"],
                "name": it["name"],
                "type": it["type"],
                "sources": set(),
                "count": 0,
            }
        agg[key]["sources"].add(it["source"])
        agg[key]["count"] += 1
        if agg[key]["count"] < agg[key]["count"]:  # pragma: no cover
            pass
    return agg


def token_set(s):
    return set(re.findall(r"[a-zа-яё0-9]+", norm(s)))


SERGEK_VARIANTS = ("патруль", "трасса", "трассы", "перекрест", "линейн", "lite", "box")


def sergek_variant(name):
    n = norm(name)
    for v in SERGEK_VARIANTS:
        if v in n:
            return v
    return None


def names_match(a, b):
    na, nb = norm(a), norm(b)
    if not na or not nb:
        return False
    if na == nb:
        return True

    va, vb = sergek_variant(a), sergek_variant(b)
    if va and vb and va != vb:
        return False

    if na in nb or nb in na:
        if min(len(na), len(nb)) >= 10:
            return True
    ta, tb = token_set(a), token_set(b)
    if not ta or not tb:
        return False
    overlap = len(ta & tb) / max(len(ta), len(tb))
    if overlap >= 0.75 and min(len(na), len(nb)) >= 15:
        return True
    aliases = [
        ("сергек патруль", "сергек. патруль"),
        ("сергек трасса", "сергек. трасса"),
        ("сергек трассы", "сергек. трасса"),
        ("сергек линейный", "сергек. линейный участок"),
        ("обслуживание лу", "обслуживание линей"),
        ("обслуживание п", "обслуживание перекрест"),
        ("интернет в цод", "интернет для цода"),
        ("интернет 100 мбит", "интернет для цода"),
        ("техническое обслуживание цод", "техническое обслуживание цода"),
        ("канал передачи данных линейные", "канал передачи данных - линейные"),
        ("канал передачи данных перекрестки", "канал передачи данных - перекрестки"),
        ("система хранения данных птб", "система хранения данных"),
        ("монтаж апк на линейном", "монтаж апк на линейном участке"),
        ("крепление и навеска", "монтаж «сергек. патруль»"),
        ("настройка пусконаладка", "монтаж и запуск серверов"),
        ("настройка пуско-наладка", "монтаж и запуск серверов"),
        ("затраты на электроэнергию", "стоимость электроэнергии"),
        ("непредвиденные работы", "непредвиденные работы"),
    ]
    for x, y in aliases:
        if (x in na and y in nb) or (x in nb and y in na):
            return True
    return False


def is_important_product(name):
    key = norm(name)
    return bool(
        re.search(
            r"сергек|патруль|трасса|lite|box|метеостанц|видеостен|"
            r"транспортный\s+двойник|ngfw|шлюз\s+безопасности|"
            r"мобильное\s+рабочее\s+место|mdm|резервный\s+сервер",
            key,
        )
    )


def find_missing(main_rows, hist_agg, min_count=2):
    missing = []
    main_names = [r["name"] for r in main_rows]

    for key, data in sorted(hist_agg.items(), key=lambda x: -x[1]["count"]):
        if data["count"] < min_count and not is_important_product(data["name"]):
            continue
        if not is_valid_artifact(data["name"]):
            continue
        if not should_add_missing(data):
            continue
        found = False
        for main_name in main_names:
            if names_match(data["name"], main_name):
                found = True
                break
        if not found:
            missing.append(data)
    return missing


def build_workbook(main_wb, main_rows, missing):
    out = openpyxl.load_workbook(MAIN)
    ws = out["Артефакты"]

    # Insert CAPEX/OPEX after column C (Наименование) => col 4
    ws.insert_cols(4)
    ws.cell(2, 4, "CAPEX/OPEX")

    # Insert source columns after existing data (after col 10 -> now shifted)
    # Original col 10 was Поставщик характеристик; after insert it's col 11
    ws.insert_cols(12, amount=2)
    ws.cell(2, 12, "Источник (исторические ФЭМ)")
    ws.cell(2, 13, "Встречаемость в истор. ФЭМ")

    all_rows = []

    # existing rows
    for i, row in enumerate(main_rows):
        r = row["row"]
        capex_opex = classify(row["category"])
        ws.cell(r, 4, capex_opex)
        all_rows.append(
            {
                "category": row["category"],
                "name": row["name"],
                "type": capex_opex,
                "unit": row["unit"],
                "price": row["price"],
                "qty_supplier": row["qty_supplier"],
                "stage": row["stage"],
                "price_supplier": row["price_supplier"],
                "calc_holder": row["calc_holder"],
                "char_supplier": row["char_supplier"],
                "source": None,
                "count": None,
            }
        )

    next_row = ws.max_row + 1
    added = []
    for data in missing:
        cat = data["category"]
        # map generic historical categories to main-file style where possible
        cat_norm = norm(cat)
        if "центр обработки" in cat_norm or data["type"] == "CAPEX" and "сервер" in norm(data["name"]):
            cat = "Центр обработки и анализа данных"
        elif data["type"] == "OPEX" and any(x in cat_norm for x in ["канал", "связ"]):
            cat = "Аренда каналов связи"
        elif data["type"] == "OPEX" and "обслуж" in cat_norm:
            cat = "Обслуживание АПК" if "апк" in norm(data["name"]) else "Услуги сторонних организаций"
        elif data["type"] == "OPEX" and any(x in norm(data["name"]) for x in ["гсм", "канц", "хоз", "охрана", "клининг", "билет", "проживание", "суточн"]):
            cat = "Экономические предположения"
        elif data["type"] == "CAPEX" and any(x in norm(data["name"]) for x in ["монтаж", "креплен", "навеск", "интеграц"]):
            cat = 'Стоимость монтажных работ и услуг для запуска проекта'
        elif "аппаратно-программ" in cat_norm or "сергек" in norm(data["name"]):
            cat = 'Аппаратно-программные комплексы "Сергек"'
        elif "камер" in cat_norm or "овн" in norm(data["name"]):
            cat = "Камеры общественной безопасности"

        ws.cell(next_row, 2, cat)
        ws.cell(next_row, 3, data["name"])
        ws.cell(next_row, 4, data["type"])
        ws.cell(next_row, 12, "; ".join(sorted(data["sources"])[:5]))
        ws.cell(next_row, 13, data["count"])

        item = {
            "category": cat,
            "name": data["name"],
            "type": data["type"],
            "unit": None,
            "price": None,
            "qty_supplier": None,
            "stage": None,
            "price_supplier": None,
            "calc_holder": None,
            "char_supplier": None,
            "source": "; ".join(sorted(data["sources"])[:5]),
            "count": data["count"],
        }
        all_rows.append(item)
        added.append(item)
        next_row += 1

    # Create CAPEX / OPEX sheets
    for sheet_name, flt in [("CAPEX", "CAPEX"), ("OPEX", "OPEX")]:
        if sheet_name in out.sheetnames:
            del out[sheet_name]
        nws = out.create_sheet(sheet_name)
        headers = [
            "Статья затрат (1-уровень)",
            "Наименование (Артефакт)",
            "CAPEX/OPEX",
            "Ед. измерения",
            "Цена за ед.с НДС, тенге",
            "Поставщик количества",
            "Этап",
            "Поставщик цены",
            "Держатель калькулятора",
            "Поставщик характеристик",
            "Источник (исторические ФЭМ)",
            "Встречаемость в истор. ФЭМ",
        ]
        for c, h in enumerate(headers, 2):
            nws.cell(2, c, h)
        rr = 3
        for item in all_rows:
            if item["type"] != flt:
                continue
            nws.cell(rr, 2, item["category"])
            nws.cell(rr, 3, item["name"])
            nws.cell(rr, 4, item["type"])
            nws.cell(rr, 5, item["unit"])
            nws.cell(rr, 6, item["price"])
            nws.cell(rr, 7, item["qty_supplier"])
            nws.cell(rr, 8, item["stage"])
            nws.cell(rr, 9, item["price_supplier"])
            nws.cell(rr, 10, item["calc_holder"])
            nws.cell(rr, 11, item["char_supplier"])
            nws.cell(rr, 12, item["source"])
            nws.cell(rr, 13, item["count"])
            rr += 1

    out.save(OUTPUT)
    return added, all_rows


def main():
    hist_items = extract_historical()
    hist_agg = aggregate_historical(hist_items)
    main_wb, main_rows = load_main_artifacts()
    missing = find_missing(main_rows, hist_agg, min_count=2)

    added, all_rows = build_workbook(main_wb, main_rows, missing)

    report = {
        "main_before": len(main_rows),
        "historical_files": len(list(HIST_DIR.glob("*.xlsx"))),
        "unique_historical_artifacts": len(hist_agg),
        "added": len(added),
        "added_capex": sum(1 for x in added if x["type"] == "CAPEX"),
        "added_opex": sum(1 for x in added if x["type"] == "OPEX"),
        "total_after": len(all_rows),
        "capex_total": sum(1 for x in all_rows if x["type"] == "CAPEX"),
        "opex_total": sum(1 for x in all_rows if x["type"] == "OPEX"),
        "missing_list": [
            {
                "Статья затрат (1-уровень)": x["category"],
                "Наименование (Артефакт)": x["name"],
                "CAPEX/OPEX": x["type"],
                "Источник (исторические ФЭМ)": x["source"],
                "Встречаемость": x["count"],
            }
            for x in added
        ],
    }
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps({k: report[k] for k in report if k != "missing_list"}, ensure_ascii=False, indent=2))
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    main()
