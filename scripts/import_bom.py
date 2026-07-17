#!/usr/bin/env python3
"""Export BOM sheets from Sergek Excel specs to JSON templates."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
INPUT_DIR = ROOT / "drive-input"
OUTPUT_DIR = ROOT / "data" / "bom-templates"

SHEET_MAP = {
    ("LU", "Hikvision"): "ЛУ_до 3 полос",
    ("P", "Hikvision"): "П",
    ("LU", "Dahua"): "ЛУ_до 3 полос",
    ("P", "Dahua"): "П",
}


def _manufacturer_file(mfr: str) -> Path:
    key = "Hikvision" if mfr == "Hikvision" else "Dahua"
    for path in INPUT_DIR.glob("*.xlsx"):
        if key in path.name:
            return path
    raise FileNotFoundError(f"No workbook for {mfr}")


def import_sheet(segment: str, manufacturer: str) -> dict:
    sheet_name = SHEET_MAP[(segment, manufacturer)]
    wb = load_workbook(_manufacturer_file(manufacturer), data_only=False)
    ws = wb[sheet_name]

    rows = []
    for r in range(2, ws.max_row + 1):
        num = ws[f"A{r}"].value
        if num is None:
            continue
        g_val = ws[f"G{r}"].value
        h_val = ws[f"H{r}"].value
        rows.append(
            {
                "row": int(num),
                "module": (ws[f"B{r}"].value or "").strip() or None,
                "description": (ws[f"C{r}"].value or "").strip(),
                "name": (ws[f"D{r}"].value or "").strip(),
                "unit": ws[f"F{r}"].value,
                "qty_pp": g_val if isinstance(g_val, str) and g_val.startswith("=") else g_val,
                "qty_light": h_val if isinstance(h_val, str) and h_val.startswith("=") else h_val,
                "qty_pp_formula": g_val if isinstance(g_val, str) and g_val.startswith("=") else None,
                "qty_light_formula": h_val if isinstance(h_val, str) and h_val.startswith("=") else None,
                "note": (ws[f"I{r}"].value or "").strip() or None,
            }
        )

    return {
        "meta": {
            "segment": segment,
            "manufacturer": manufacturer,
            "apk_version": "4.0",
            "source_sheet": sheet_name,
        },
        "rows": rows,
    }


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for segment, manufacturer in SHEET_MAP:
        data = import_sheet(segment, manufacturer)
        out = OUTPUT_DIR / f"{segment.lower()}-{manufacturer.lower()}.json"
        out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"Wrote {out} ({len(data['rows'])} rows)")


if __name__ == "__main__":
    main()
