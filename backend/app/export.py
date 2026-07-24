"""Export consolidated specification to Excel."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from app.models import ProjectSummary


def export_to_excel(summary: ProjectSummary) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводная"

    headers = ["№", "Описание", "Наименование", "Ед.", "Кол-во", "Цена, ₸", "Сумма, ₸"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in summary.consolidated:
        ws.append(
            [
                row.get("row"),
                row.get("description"),
                row.get("name"),
                row.get("unit"),
                row.get("total_qty"),
                row.get("unit_price_kzt"),
                row.get("total_price_kzt"),
            ]
        )

    ws.append([])
    ws.append(["", "", "", "", "", "Итого:", summary.total_cost_kzt])

    # Breakdown sheet
    bd = wb.create_sheet("Разбивка")
    bd.append(["Конфигурация", "Сегмент", "Питание", "АПК", "Позиций", "Сумма, ₸"])
    for cfg in summary.configurations:
        bd.append(
            [
                cfg["label"],
                cfg["segment"],
                cfg["power_type"],
                cfg["apk_count"],
                cfg["line_count"],
                cfg.get("subtotal_kzt"),
            ]
        )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
